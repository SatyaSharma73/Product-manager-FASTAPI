import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── colour palette ──────────────────────────────────────
C_NAV    = '1E1B4B'
C_WHITE  = 'FFFFFF'
C_CRIT   = 'FEE2E2'
C_CRIT_T = 'DC2626'
C_CRIT_B = 'FECACA'
C_MED    = 'FEF9C3'
C_MED_T  = 'B45309'
C_MED_B  = 'FDE68A'
C_OK     = 'F0FDF4'
C_OK_T   = '16A34A'
C_OK_B   = 'BBF7D0'
C_STRIPE = 'F8F7FF'
C_TITLE  = '4F46E5'

def fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)

def font(bold=False, color=C_NAV, size=10, name='Calibri', italic=False):
    return Font(bold=bold, color=color, size=size, name=name, italic=italic)

def thin_border(color='D1D5DB'):
    s = Side(style='thin', color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def align(h='left', v='center', wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

# ════════════════════════════════════════════════════════
# SHEET 1 — Bug Fix Log
# ════════════════════════════════════════════════════════
ws = wb.active
ws.title = 'Bug Fix Log'

# title
ws.merge_cells('A1:H1')
t = ws['A1']
t.value     = 'Product Manager App — Bug Fix Log'
t.font      = Font(bold=True, color=C_WHITE, size=14, name='Calibri')
t.fill      = fill(C_NAV)
t.alignment = align('center')
ws.row_dimensions[1].height = 36

# subtitle
ws.merge_cells('A2:H2')
s = ws['A2']
s.value     = 'All confirmed bugs identified during code audit — root cause + fix documented'
s.font      = Font(color='A5B4FC', size=10, name='Calibri', italic=True)
s.fill      = fill(C_NAV)
s.alignment = align('center')
ws.row_dimensions[2].height = 20

# column headers
headers    = ['#', 'File', 'Severity', 'Bug Title', 'What Was Wrong', 'Root Cause', 'Fix Applied', 'Status']
col_widths = [4,    22,     12,          30,           52,               44,            52,             12]

for ci, (h, w) in enumerate(zip(headers, col_widths), start=1):
    cell = ws.cell(row=3, column=ci, value=h)
    cell.font      = Font(bold=True, color=C_WHITE, size=10, name='Calibri')
    cell.fill      = fill(C_TITLE)
    cell.alignment = align('center')
    cell.border    = thin_border(C_TITLE)
    ws.column_dimensions[get_column_letter(ci)].width = w

ws.row_dimensions[3].height = 22

# ── bug records ─────────────────────────────────────────
bugs = [
    {
        'id': 1,
        'file': 'EditModal.js',
        'sev': ('Critical', C_CRIT, C_CRIT_T, C_CRIT_B),
        'title': 'Save button stuck in loading state after success',
        'what': (
            'After successfully saving an edited product, the Save Changes button '
            'remained in a spinning / disabled loading state. The modal could not be '
            'reused without closing and reopening it.'
        ),
        'cause': (
            'setLoading(false) was only called inside the catch block. On a successful '
            'save, onSaved() was called directly with no cleanup, leaving loading=true permanently.'
        ),
        'fix': (
            'Moved setLoading(false) from the catch block into a finally block. '
            'The finally block runs on both success and failure, ensuring the button '
            'always resets regardless of the outcome.'
        ),
    },
    {
        'id': 2,
        'file': 'DeleteModal.js',
        'sev': ('Critical', C_CRIT, C_CRIT_T, C_CRIT_B),
        'title': 'Delete button stuck in loading state after success',
        'what': (
            'After successfully deleting a product, the Delete button remained spinning '
            'and disabled. The modal appeared frozen after a successful delete, making '
            'it impossible to delete another product without a page refresh.'
        ),
        'cause': (
            'Identical root cause to EditModal.js: setLoading(false) was only in the '
            'catch block, so a successful deletion never reset the loading state.'
        ),
        'fix': (
            'Moved setLoading(false) to a finally block so it always executes after '
            'the async operation, regardless of whether it succeeded or failed.'
        ),
    },
    {
        'id': 3,
        'file': 'FindByIdTab.js',
        'sev': ('Medium', C_MED, C_MED_T, C_MED_B),
        'title': 'Unhelpful generic error for invalid UUID input',
        'what': (
            'Entering non-UUID text (e.g. a product name or random string) returned '
            'a generic "Something went wrong" message. A 422 Unprocessable Entity '
            'response from FastAPI was not handled separately from server errors.'
        ),
        'cause': (
            'The catch block only checked for HTTP 404 and fell through to a generic '
            'message for everything else, including 422 (invalid UUID format) and '
            'network failures. No client-side UUID validation existed before the request.'
        ),
        'fix': (
            'Added a UUID regex check before the API call — invalid format is caught '
            'instantly without a network round-trip. Added specific messages for: '
            '422 (invalid UUID), 404 (not found), and server unreachable. '
            'Regex: /^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$/i'
        ),
    },
    {
        'id': 4,
        'file': 'ImportTab.js',
        'sev': ('Medium', C_MED, C_MED_T, C_MED_B),
        'title': 'Empty Excel file silently shows blank preview table',
        'what': (
            'Uploading an Excel file that contained only a header row (no product data) '
            'silently advanced to the preview step, displaying a completely empty table '
            'with no error, warning, or guidance shown to the user.'
        ),
        'cause': (
            'After the /previewimport API returned successfully, the code unconditionally '
            'moved to the PREVIEW step. It never checked whether the returned rows '
            'array was empty before proceeding to render the preview table.'
        ),
        'fix': (
            'Added a check immediately after the API response: '
            'if (data.rows.length === 0) show an error and stop. '
            'User now sees: "No data rows found in this file. Make sure there is at '
            'least one product row below the header." and stays on the drop zone step.'
        ),
    },
    {
        'id': 5,
        'file': 'ProductsTab.js',
        'sev': ('Medium', C_MED, C_MED_T, C_MED_B),
        'title': 'Non-numeric price filter silently hides all products',
        'what': (
            'Typing letters or symbols (e.g. "abc") into the Min Price or Max Price '
            'filter fields caused all products to vanish from the table with no error '
            'shown. The filter appeared active but returned zero results.'
        ),
        'cause': (
            'parseFloat("abc") returns NaN in JavaScript. Any comparison involving NaN '
            '(e.g. price < NaN) always evaluates to false, so no product passed the '
            'filter condition — silently hiding everything without any feedback.'
        ),
        'fix': (
            'Added !isNaN(parseFloat(value)) guard before each price comparison. '
            'The filter now only applies when the value is an actual valid number. '
            'Non-numeric input is ignored and all products continue to show normally.'
        ),
    },
]

for r_idx, bug in enumerate(bugs, start=4):
    row_bg = C_STRIPE if r_idx % 2 == 0 else C_WHITE
    sev_label, sev_bg, sev_text, sev_bdr = bug['sev']
    values = [bug['id'], bug['file'], sev_label, bug['title'],
              bug['what'], bug['cause'], bug['fix'], 'Fixed ✓']

    for ci, val in enumerate(values, start=1):
        cell = ws.cell(row=r_idx, column=ci, value=val)
        cell.border    = thin_border()
        cell.alignment = align('center' if ci in (1, 3, 8) else 'left', wrap=True)
        cell.font      = Font(size=9, name='Calibri', color=C_NAV)
        if ci not in (3, 8):
            cell.fill = fill(row_bg)

    # severity badge
    sc = ws.cell(row=r_idx, column=3)
    sc.font   = Font(bold=True, color=sev_text, size=9, name='Calibri')
    sc.fill   = fill(sev_bg)
    sc.border = thin_border(sev_bdr)

    # status badge
    stc = ws.cell(row=r_idx, column=8)
    stc.font   = Font(bold=True, color=C_OK_T, size=9, name='Calibri')
    stc.fill   = fill(C_OK)
    stc.border = thin_border(C_OK_B)

    # dynamic row height
    longest = max(len(bug['what']), len(bug['cause']), len(bug['fix']))
    ws.row_dimensions[r_idx].height = max(70, min(130, longest // 2))

# freeze + autofilter
ws.freeze_panes = 'A4'
ws.auto_filter.ref = f'A3:H{3 + len(bugs)}'

# ════════════════════════════════════════════════════════
# SHEET 2 — Summary
# ════════════════════════════════════════════════════════
dash = wb.create_sheet('Summary')

dash.merge_cells('A1:C1')
dt = dash['A1']
dt.value     = 'Bug Fix Summary'
dt.font      = Font(bold=True, color=C_WHITE, size=13, name='Calibri')
dt.fill      = fill(C_NAV)
dt.alignment = align('center')
dash.row_dimensions[1].height = 32

# spacer
dash.row_dimensions[2].height = 8

# stats block
stats = [
    ('Total Bugs Found', 5, C_TITLE, 'EEF2FF', 'C7D2FE'),
    ('Critical Bugs',    2, C_CRIT_T, C_CRIT,  C_CRIT_B),
    ('Medium Bugs',      3, C_MED_T,  C_MED,   C_MED_B),
    ('All Fixed',        5, C_OK_T,   C_OK,    C_OK_B),
]
for ri, (label, count, tc, bg, bdr) in enumerate(stats, start=3):
    lc = dash.cell(row=ri, column=1, value=label)
    lc.font      = Font(bold=True, color=tc, size=10, name='Calibri')
    lc.fill      = fill(bg)
    lc.border    = thin_border(bdr)
    lc.alignment = align('left')

    vc = dash.cell(row=ri, column=2, value=count)
    vc.font      = Font(bold=True, color=tc, size=12, name='Calibri')
    vc.fill      = fill(bg)
    vc.border    = thin_border(bdr)
    vc.alignment = align('center')
    dash.row_dimensions[ri].height = 26

# spacer
dash.row_dimensions[7].height = 14

# file breakdown header
for ci, h in enumerate(['File', 'Bugs Fixed', 'Severity'], start=1):
    hc = dash.cell(row=8, column=ci, value=h)
    hc.font      = Font(bold=True, color=C_WHITE, size=10, name='Calibri')
    hc.fill      = fill(C_TITLE)
    hc.alignment = align('center')
    hc.border    = thin_border(C_TITLE)
dash.row_dimensions[8].height = 22

file_rows = [
    ('EditModal.js',    1, 'Critical'),
    ('DeleteModal.js',  1, 'Critical'),
    ('FindByIdTab.js',  1, 'Medium'),
    ('ImportTab.js',    1, 'Medium'),
    ('ProductsTab.js',  1, 'Medium'),
]
for ri, (fname, cnt, sev) in enumerate(file_rows, start=9):
    bg = C_STRIPE if ri % 2 == 0 else C_WHITE
    fc = dash.cell(row=ri, column=1, value=fname)
    fc.font = Font(size=10, name='Calibri', color=C_NAV)
    fc.fill = fill(bg); fc.border = thin_border(); fc.alignment = align('left')

    cc = dash.cell(row=ri, column=2, value=cnt)
    cc.font = Font(bold=True, size=10, name='Calibri', color=C_OK_T)
    cc.fill = fill(C_OK); cc.border = thin_border(C_OK_B); cc.alignment = align('center')

    sev_bg   = C_CRIT   if sev == 'Critical' else C_MED
    sev_text = C_CRIT_T if sev == 'Critical' else C_MED_T
    sev_bdr  = C_CRIT_B if sev == 'Critical' else C_MED_B
    sc = dash.cell(row=ri, column=3, value=sev)
    sc.font = Font(bold=True, size=10, name='Calibri', color=sev_text)
    sc.fill = fill(sev_bg); sc.border = thin_border(sev_bdr); sc.alignment = align('center')
    dash.row_dimensions[ri].height = 20

dash.column_dimensions['A'].width = 22
dash.column_dimensions['B'].width = 14
dash.column_dimensions['C'].width = 14

wb.save(r'C:\Users\2153940\OneDrive - Cognizant\Desktop\bug_fix_log.xlsx')
print('Done')
